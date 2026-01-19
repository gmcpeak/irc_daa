# worker/worker.py
import os
import io
import json
import logging
import re
import time
import uuid
import queue
import threading
from datetime import datetime
from typing import Optional, Dict, List, Tuple

import dotenv
from fastapi import FastAPI, UploadFile, File, HTTPException, Query
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, PlainTextResponse, StreamingResponse
from pydantic import BaseModel
from PIL import Image, ImageOps
from PIL.Image import Image as PILImage
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openai import OpenAI

from box_sdk_gen import (
    BoxClient,
    BoxJWTAuth,
    FileWithInMemoryCacheTokenStorage,
    JWTConfig,
)
from box_sdk_gen.managers.uploads import (
    UploadFileAttributes,
    UploadFileAttributesParentField,
)

from base64 import b64encode

# -------------------------------------------------
# basic setup
# -------------------------------------------------
dotenv.load_dotenv(".env")  # uses your file baked into the image

logging.basicConfig(level=logging.INFO)
logging.getLogger("box_sdk_gen").setLevel(logging.CRITICAL)

PROMPT_PATH = "./assets/system_prompt.txt"
DOC1_PATH = "./assets/Goodwill-Donation-Value-Guide.txt"
DOC2_PATH = "./assets/Salvation-Army-Donation-Value-Guide.txt"

OPENAI_KEY = os.getenv("OPENAI_KEY")
if not OPENAI_KEY:
    raise RuntimeError("OPENAI_KEY not configured in .env")

# Box upload target folder
BOX_FOLDER_ID = os.getenv("BOX_FOLDER_ID", "333439130395")

# Job retention (in-memory). If the container restarts, jobs are lost.
JOB_TTL_SECONDS = int(os.getenv("JOB_TTL_SECONDS", "3600"))  # 1 hour
MAX_TEXT_CHARS_STORED = int(os.getenv("MAX_TEXT_CHARS_STORED", "200000"))

client = OpenAI(api_key=OPENAI_KEY)

LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

app = FastAPI(title="Donation Audit Worker")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

# -------------------------------------------------
# job model + store
# -------------------------------------------------
class JobStatus:
    QUEUED = "queued"
    RUNNING = "running"
    DONE = "done"
    FAILED = "failed"


class JobRecord(BaseModel):
    job_id: str
    status: str
    expected_filename: str

    created_at: float
    started_at: Optional[float] = None
    finished_at: Optional[float] = None

    error: Optional[str] = None

    # outputs
    output_text: Optional[str] = None
    docx_bytes: Optional[bytes] = None

    # Box metadata
    box_file_id: Optional[str] = None
    box_file_name: Optional[str] = None


jobs_lock = threading.Lock()
jobs: Dict[str, JobRecord] = {}

# -------------------------------------------------
# in-memory queue + worker thread
# -------------------------------------------------
job_queue: "queue.Queue[Tuple[str, bytes, str]]" = queue.Queue()


def _read_texts(paths: List[str]) -> str:
    parts: List[str] = []
    for p in paths:
        try:
            with open(p, "r", encoding="utf-8", errors="ignore") as f:
                parts.append(f.read())
        except Exception:
            logging.exception("Failed reading %s", p)
    return "\n".join(parts)


def _load_pil_from_bytes(data: bytes) -> Optional[PILImage]:
    if not data:
        return None
    try:
        img = Image.open(io.BytesIO(data))
        return ImageOps.exif_transpose(img)
    except Exception:
        return None


def _openai_process(prompt_text: str, image: Optional[PILImage]) -> str:
    """
    Returns plain text. Note: we do NOT enable web_search tools here to avoid
    special citation tokens being inserted into the output.
    """
    content = []
    if prompt_text:
        content.append({"type": "input_text", "text": prompt_text})
    if isinstance(image, PILImage):
        buf = io.BytesIO()
        image.save(buf, format="PNG")
        b64 = b64encode(buf.getvalue()).decode("utf-8")
        content.append(
            {"type": "input_image", "image_url": f"data:image/png;base64,{b64}"}
        )

    resp = client.responses.create(
        model=os.getenv("OPENAI_MODEL", "gpt-5.1"),
        reasoning={"effort": "low"},
        text={"verbosity": "low"},
        # tools intentionally omitted to keep output clean
        input=[{"role": "user", "content": content}] if content else prompt_text,
    )

    text = getattr(resp, "output_text", None)
    if not text:
        try:
            text = resp.output[0].content[0].text
        except Exception:
            text = str(resp)
    return text or ""


def add_hyperlink(paragraph, text, url):
    """
    Add a hyperlink to a paragraph. Returns the hyperlink run.
    """
    part = paragraph.part
    r_id = part.relate_to(
        url,
        reltype=(
            "http://schemas.openxmlformats.org/officeDocument/2006/"
            "relationships/hyperlink"
        ),
        is_external=True,
    )

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")

    rPr = OxmlElement("w:rPr")

    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)

    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0000FF")
    rPr.append(color)

    new_run.append(rPr)

    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)

    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)

    return new_run


def _docx_bytes_from_text(
    text: str,
    image: Optional[PILImage],
    title: str = "Donation Audit",
) -> bytes:
    doc = Document()
    if title:
        doc.add_heading(title, level=0)

    # optional image at top
    if image is not None:
        img_buf = io.BytesIO()
        image.save(img_buf, format="PNG")
        img_buf.seek(0)
        section = doc.sections[0]
        max_width = (
            section.page_width - section.left_margin - section.right_margin
        )
        doc.add_picture(img_buf, width=max_width)
        doc.add_paragraph("")

    # Add text (convert markdown links to docx hyperlinks)
    for line in (text or "").splitlines():
        p = doc.add_paragraph()
        pos = 0
        for match in LINK_RE.finditer(line):
            start, end = match.span()
            link_text = match.group(1)
            link_url = match.group(2)

            if start > pos:
                p.add_run(line[pos:start])

            add_hyperlink(p, link_text, link_url)
            pos = end

        if pos < len(line):
            p.add_run(line[pos:])

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def get_box_client_from_config(config_path: str = "config.json") -> BoxClient:
    with open(config_path, "r", encoding="utf-8") as f:
        cfg_json = json.load(f)

    user_id = cfg_json.get("sub")

    jwt_cfg = JWTConfig.from_config_file(
        config_file_path=config_path,
        token_storage=FileWithInMemoryCacheTokenStorage(".user.jwt.cache"),
    )

    auth = BoxJWTAuth(jwt_cfg)
    if user_id:
        auth = auth.with_user_subject(user_id)

    return BoxClient(auth)


def upload_docx_bytes(
    client: BoxClient,
    file_name: str,
    docx_bytes: bytes,
    folder_id: str,
):
    stream = io.BytesIO(docx_bytes)
    attrs = UploadFileAttributes(
        name=file_name, parent=UploadFileAttributesParentField(id=folder_id)
    )
    try:
        res = client.uploads.upload_file(attrs, stream)
        uploaded = res.entries[0]
        return uploaded.id, uploaded.name
    except Exception as err:
        from box_sdk_gen import BoxAPIError

        if isinstance(err, BoxAPIError) and err.response_info.body.get("code") == "item_name_in_use":
            conflicts = err.response_info.body["context_info"]["conflicts"]
            box_file_id = conflicts["id"]
            stream.seek(0)
            res = client.uploads.upload_file_version(box_file_id, stream, attrs)
            uploaded = res.entries[0]
            return uploaded.id, uploaded.name
        raise


def _set_job_fields(job_id: str, **fields):
    with jobs_lock:
        rec = jobs.get(job_id)
        if not rec:
            return
        updated = rec.model_copy(update=fields)
        jobs[job_id] = updated


def _process_job(job_id: str, image_bytes: bytes, expected_filename: str) -> None:
    try:
        _set_job_fields(job_id, status=JobStatus.RUNNING, started_at=time.time(), error=None)
        logging.info("BG: job_id=%s starting processing...", job_id)

        image = _load_pil_from_bytes(image_bytes)
        if image is None:
            raise RuntimeError("Could not decode image")

        system_prompt = _read_texts([PROMPT_PATH])
        guides_text = _read_texts([DOC1_PATH, DOC2_PATH])
        full_prompt = (system_prompt or "") + "\n\n" + (guides_text or "")

        logging.info("BG: job_id=%s calling OpenAI...", job_id)
        out_text = _openai_process(full_prompt, image)

        # store text early (truncate to avoid unbounded memory)
        if out_text and len(out_text) > MAX_TEXT_CHARS_STORED:
            out_text = out_text[:MAX_TEXT_CHARS_STORED] + "\n\n[TRUNCATED]"
        _set_job_fields(job_id, output_text=out_text)

        logging.info("BG: job_id=%s building DOCX...", job_id)
        docx_bytes = _docx_bytes_from_text(out_text or "", image)

        # Box upload
        logging.info("BG: job_id=%s uploading DOCX to Box...", job_id)
        box_client = get_box_client_from_config("config.json")
        box_file_id, box_file_name = upload_docx_bytes(
            box_client, expected_filename, docx_bytes, folder_id=BOX_FOLDER_ID
        )

        _set_job_fields(
            job_id,
            status=JobStatus.DONE,
            finished_at=time.time(),
            docx_bytes=docx_bytes,
            box_file_id=str(box_file_id),
            box_file_name=str(box_file_name),
        )
        logging.info("BG: job_id=%s done. Box id=%s", job_id, box_file_id)

    except Exception as e:
        logging.exception("BG: job_id=%s failed.", job_id)
        _set_job_fields(
            job_id,
            status=JobStatus.FAILED,
            finished_at=time.time(),
            error=str(e),
        )


def _queue_worker_loop():
    logging.info("QueueWorker: starting background queue worker loop...")
    while True:
        try:
            job_id, image_bytes, expected_filename = job_queue.get()
            logging.info("QueueWorker: picked up job_id=%s (queue size=%d)", job_id, job_queue.qsize())
            _process_job(job_id, image_bytes, expected_filename)
        except Exception:
            logging.exception("QueueWorker: unexpected error.")
        finally:
            job_queue.task_done()


def _cleanup_loop():
    while True:
        time.sleep(30)
        now = time.time()
        with jobs_lock:
            to_delete = []
            for jid, rec in jobs.items():
                age = now - rec.created_at
                if age > JOB_TTL_SECONDS:
                    to_delete.append(jid)
            for jid in to_delete:
                jobs.pop(jid, None)


_worker_thread = threading.Thread(target=_queue_worker_loop, name="queue-worker-thread", daemon=True)
_worker_thread.start()

_cleanup_thread = threading.Thread(target=_cleanup_loop, name="jobs-cleanup-thread", daemon=True)
_cleanup_thread.start()


# -------------------------------------------------
# API schema + endpoints
# -------------------------------------------------
class ProcessResponse(BaseModel):
    status_message: str
    expected_filename: str
    job_id: str


class JobStatusResponse(BaseModel):
    job_id: str
    status: str
    expected_filename: str
    created_at: float
    started_at: Optional[float] = None
    finished_at: Optional[float] = None
    error: Optional[str] = None
    download_ready: bool
    output_text_ready: bool
    box_file_id: Optional[str] = None
    box_file_name: Optional[str] = None
    output_text: Optional[str] = None


@app.get("/health")
def health():
    return {"ok": True}


@app.post("/process", response_model=ProcessResponse)
async def process_endpoint(file: UploadFile = File(...)):
    raw_bytes = await file.read()
    image = _load_pil_from_bytes(raw_bytes)
    if image is None:
        raise HTTPException(status_code=400, detail="Could not decode image")

    job_id = uuid.uuid4().hex
    expected_filename = (
        f"GW-IRC-Donation-Value-Audit--{datetime.now().strftime('%Y%m%d-%H%M%S')}.docx"
    )

    rec = JobRecord(
        job_id=job_id,
        status=JobStatus.QUEUED,
        expected_filename=expected_filename,
        created_at=time.time(),
    )
    with jobs_lock:
        jobs[job_id] = rec

    job_queue.put((job_id, raw_bytes, expected_filename))

    status_message = (
        "The worker received your image and queued processing. "
        "You may keep this tab open to download the result when ready, "
        "or retrieve it later from Box under the filename shown."
    )

    return ProcessResponse(
        status_message=status_message,
        expected_filename=expected_filename,
        job_id=job_id,
    )


@app.get("/jobs/{job_id}", response_model=JobStatusResponse)
def get_job_status(
    job_id: str,
    include_text: bool = Query(False, description="If true, include output_text in response."),
    text_truncate: int = Query(20000, ge=0, le=200000, description="Max chars of output_text to return."),
):
    with jobs_lock:
        rec = jobs.get(job_id)
    if not rec:
        raise HTTPException(status_code=404, detail="Job not found")

    download_ready = bool(rec.docx_bytes) and rec.status == JobStatus.DONE
    output_text_ready = bool(rec.output_text) and rec.status in (JobStatus.DONE, JobStatus.FAILED)

    out_text = None
    if include_text and rec.output_text:
        out_text = rec.output_text[:text_truncate] if text_truncate else rec.output_text

    return JobStatusResponse(
        job_id=rec.job_id,
        status=rec.status,
        expected_filename=rec.expected_filename,
        created_at=rec.created_at,
        started_at=rec.started_at,
        finished_at=rec.finished_at,
        error=rec.error,
        download_ready=download_ready,
        output_text_ready=output_text_ready,
        box_file_id=rec.box_file_id,
        box_file_name=rec.box_file_name,
        output_text=out_text,
    )


@app.get("/jobs/{job_id}/text")
def get_job_text(job_id: str):
    with jobs_lock:
        rec = jobs.get(job_id)
    if not rec:
        raise HTTPException(status_code=404, detail="Job not found")

    if not rec.output_text and rec.status not in (JobStatus.DONE, JobStatus.FAILED):
        # still running/queued
        return JSONResponse(status_code=202, content={"detail": "Not ready yet"})

    if rec.status == JobStatus.FAILED and not rec.output_text:
        # failed before producing any text
        raise HTTPException(status_code=500, detail=rec.error or "Job failed")

    return PlainTextResponse(rec.output_text or "")


@app.get("/jobs/{job_id}/docx")
def download_docx(job_id: str):
    with jobs_lock:
        rec = jobs.get(job_id)
    if not rec:
        raise HTTPException(status_code=404, detail="Job not found")

    if rec.status != JobStatus.DONE or not rec.docx_bytes:
        return JSONResponse(status_code=202, content={"detail": "Not ready yet"})

    filename = rec.expected_filename or "Donation-Audit.docx"
    stream = io.BytesIO(rec.docx_bytes)

    headers = {"Content-Disposition": f'attachment; filename="{filename}"'}
    return StreamingResponse(stream, media_type=DOCX_MIME, headers=headers)
